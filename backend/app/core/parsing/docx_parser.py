from dataclasses import dataclass, field

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

# En dessous de ce poids, une image est presque toujours une icone/puce
# decorative plutot qu'un contenu porteur d'info (photo, graphique, scan).
_MIN_MEANINGFUL_IMAGE_BYTES = 3000


@dataclass
class SectionUnit:
    section_title: str
    text: str
    images: list[bytes] = field(default_factory=list)


def _iter_block_items(doc):
    """Parcourt le corps du document dans l'ordre réel, paragraphes ET
    tableaux confondus (l'API haut niveau python-docx les expose séparément
    via doc.paragraphs / doc.tables, sans préserver leur position relative)."""
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _table_to_text(table: Table) -> str:
    lines = []
    for row in table.rows:
        cells = [_cell_text(cell) for cell in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _cell_text(cell: _Cell) -> str:
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def _paragraph_images(paragraph: Paragraph, doc) -> list[bytes]:
    """Octets bruts des images inserees dans ce paragraphe (balises a:blip)."""
    images = []
    for blip in paragraph._p.findall(f".//{qn('a:blip')}"):
        rid = blip.get(qn("r:embed"))
        if rid and rid in doc.part.rels:
            blob = doc.part.related_parts[rid].blob
            if len(blob) >= _MIN_MEANINGFUL_IMAGE_BYTES:
                images.append(blob)
    return images


def parse_docx(path: str) -> list[SectionUnit]:
    doc = DocxDocument(path)
    sections: list[SectionUnit] = []
    current_title = "Introduction"
    current_blocks: list[str] = []
    current_images: list[bytes] = []

    def flush() -> None:
        text = "\n".join(b for b in current_blocks if b.strip())
        if text.strip() or current_images:
            sections.append(
                SectionUnit(
                    section_title=current_title, text=text, images=list(current_images)
                )
            )

    for block in _iter_block_items(doc):
        if isinstance(block, Table):
            table_text = _table_to_text(block)
            if table_text:
                current_blocks.append(table_text)
            continue

        current_images.extend(_paragraph_images(block, doc))

        style_name = (block.style.name if block.style else "") or ""
        is_heading = style_name.startswith("Heading") or style_name == "Title"
        if is_heading and block.text.strip():
            flush()
            current_title = block.text.strip()
            current_blocks = []
            current_images = []
        else:
            current_blocks.append(block.text)

    flush()

    if not sections:
        sections.append(SectionUnit(section_title="Document", text=""))
    return sections
